from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_DIR = Path("legal")
FONT = "Times New Roman"
BODY_SIZE = 11
NOTE_SIZE = 9.5
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(95, 95, 95)
YELLOW = "FFF2CC"
ACCENT = "5B6573"


def set_run_font(run, size=BODY_SIZE, bold=None, italic=None, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def shade_run(run, fill=YELLOW):
    rpr = run._element.get_or_add_rPr()
    shd = rpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        rpr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_field(paragraph, instruction: str):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._element.extend([begin, instr, separate, display, end])
    set_run_font(run, size=9, color=GRAY)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_document(title: str, short_title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.first_line_indent = Cm(1.25)

    for name, size, before, after in (
        ("Heading 1", 13.5, 11, 6),
        ("Heading 2", 12, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(BODY_SIZE)
        style.paragraph_format.left_indent = Cm(1.15)
        style.paragraph_format.first_line_indent = Cm(-0.55)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.1

    if "Legal Note" not in [s.name for s in styles]:
        note = styles.add_style("Legal Note", WD_STYLE_TYPE.PARAGRAPH)
        note.font.name = FONT
        note._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        note._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        note.font.size = Pt(NOTE_SIZE)
        note.font.italic = True
        note.font.color.rgb = RGBColor(80, 80, 80)
        note.paragraph_format.first_line_indent = Cm(0)
        note.paragraph_format.space_after = Pt(6)
        note.paragraph_format.line_spacing = 1.1

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("ИСКРА  •  ЮРИДИЧЕСКИЕ ДОКУМЕНТЫ"), size=8.5, bold=True, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("Страница "), size=9, color=GRAY)
    add_field(fp, "PAGE")

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.first_line_indent = Cm(0)
    title_p.paragraph_format.space_before = Pt(6)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.keep_with_next = True
    set_run_font(title_p.add_run(title.upper()), size=16, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Cm(0)
    subtitle.paragraph_format.space_after = Pt(9)
    set_run_font(subtitle.add_run("чат-бот «Искра» в мессенджере MAX"), size=10.5, italic=True, color=GRAY)

    callout = doc.add_table(rows=1, cols=1)
    callout.autofit = False
    callout.columns[0].width = Cm(17)
    cell = callout.cell(0, 0)
    cell.width = Cm(17)
    set_cell_margins(cell)
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), YELLOW)
    tc_pr.append(shd)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), "D6B656")
        borders.append(edge)
    tc_pr.append(borders)
    cp = cell.paragraphs[0]
    cp.paragraph_format.first_line_indent = Cm(0)
    cp.paragraph_format.space_after = Pt(0)
    set_run_font(cp.add_run("ПРОЕКТ. "), size=9.5, bold=True)
    set_run_font(
        cp.add_run(
            "Перед публикацией заполнить все жёлтые поля, согласовать документ с юристом и проверить, "
            "что фактическая работа Чат-бота соответствует тексту."
        ),
        size=9.5,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    doc.core_properties.title = title
    doc.core_properties.subject = f"Юридический документ для чат-бота «Искра»: {short_title}"
    doc.core_properties.author = "Проект «Искра»"
    doc.core_properties.keywords = "Искра, MAX, чат-бот, персональные данные"
    return doc


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None, note=False):
    p = doc.add_paragraph(style="Legal Note" if note else None)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True, size=NOTE_SIZE if note else BODY_SIZE, color=GRAY if note else BLACK)
        set_run_font(p.add_run(text[len(bold_prefix):]), size=NOTE_SIZE if note else BODY_SIZE, italic=note, color=GRAY if note else BLACK)
    else:
        set_run_font(p.add_run(text), size=NOTE_SIZE if note else BODY_SIZE, italic=note, color=GRAY if note else BLACK)
    return p


def add_placeholder(doc: Document, label: str, value: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    set_run_font(p.add_run(f"{label}: "), bold=True, size=BODY_SIZE)
    r = set_run_font(p.add_run(value), bold=True, size=BODY_SIZE)
    shade_run(r)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.first_line_indent = Cm(-0.55)
    set_run_font(p.add_run(text), size=BODY_SIZE)
    return p


def add_heading(doc: Document, text: str, level=1, *, page_break_before=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.page_break_before = page_break_before
    for run in p.runs:
        set_run_font(run, size=13.5 if level == 1 else 12, bold=True)
    return p


def add_definition(doc: Document, term: str, definition: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    set_run_font(p.add_run(f"{term} — "), bold=True, size=BODY_SIZE)
    set_run_font(p.add_run(definition), size=BODY_SIZE)
    return p


def add_signature_block(doc: Document):
    add_heading(doc, "Реквизиты и контакты", 1)
    add_placeholder(doc, "Оператор/владелец", "[ПОЛНОЕ НАИМЕНОВАНИЕ ИЛИ ФИО]")
    add_placeholder(doc, "ИНН", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "ОГРН/ОГРНИП", "[ЗАПОЛНИТЬ, ЕСЛИ ПРИМЕНИМО]")
    add_placeholder(doc, "Адрес", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "Электронная почта", "[ЗАПОЛНИТЬ ПОСЛЕ СОЗДАНИЯ]")
    add_placeholder(doc, "Телефон поддержки", "[ЗАПОЛНИТЬ ИЛИ УДАЛИТЬ ПУНКТ]")
    add_placeholder(doc, "Дата вступления в силу", "[ДД.ММ.ГГГГ]")


def build_agreement() -> Path:
    doc = configure_document("Пользовательское соглашение", "пользовательское соглашение")
    add_body(
        doc,
        "Настоящее Пользовательское соглашение является публичной офертой и определяет условия "
        "использования чат-бота «Искра», размещённого в мессенджере MAX (далее — «Чат-бот»).",
    )

    add_heading(doc, "1. Основные понятия")
    add_definition(doc, "Оператор", "лицо, владеющее и управляющее Чат-ботом; его сведения приведены в конце Соглашения.")
    add_definition(doc, "Обращение", "сообщение о проблеме, направленное через Чат-бот вместе с предоставленными сведениями и материалами.")

    add_heading(doc, "2. Назначение Чат-бота")
    add_body(
        doc,
        "2.1. Чат-бот предназначен для автоматизации, систематизации и направления обращений граждан "
        "уполномоченным сотрудникам, государственным органам и органам местного самоуправления, "
        "компетентным рассматривать соответствующее обращение.",
    )
    add_body(doc, "2.2. Чат-бот позволяет создать обращение, приложить необходимые материалы, получать уведомления о ходе его обработки и получить ответ.")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    r = set_run_font(p.add_run("[ЮРИСТУ: определить, признаётся ли сообщение через Чат-бот официальным обращением в смысле Федерального закона № 59-ФЗ и иных применимых норм.]"), bold=True)
    shade_run(r)
    add_body(doc, "2.3. Чат-бот не является службой экстренной помощи. При непосредственной угрозе жизни, здоровью или безопасности следует обратиться по номеру 112 либо в соответствующую экстренную службу.")

    add_heading(doc, "3. Принятие Соглашения")
    add_body(doc, "3.1. До начала использования Чат-бота Пользователь знакомится с настоящим Соглашением и Политикой обработки персональных данных.")
    add_body(doc, "3.2. Нажатие отдельной кнопки «Принимаю пользовательское соглашение» означает полное и безоговорочное принятие Соглашения. Если Пользователь не согласен с его условиями, он не должен использовать Чат-бот.")
    add_body(doc, "3.3. Согласие на обработку персональных данных запрашивается и фиксируется отдельно от принятия настоящего Соглашения.")

    add_heading(doc, "4. Создание и рассмотрение обращения")
    add_body(doc, "4.1. Для создания обращения Пользователь предоставляет:")
    for item in (
        "фамилию, имя и отчество (при наличии);",
        "контактный номер телефона;",
        "муниципальный округ, на территории которого находится указанная в обращении проблема;",
        "текст обращения;",
        "фотографии, геолокацию и/или точный адрес проблемы — при наличии и необходимости.",
    ):
        add_bullet(doc, item)
    add_body(doc, "4.2. Пользователь обязан предоставлять достоверные сведения, относящиеся к существу обращения, и по возможности не указывать избыточные персональные данные свои или третьих лиц.")
    add_body(doc, "4.3. Пользователь подтверждает, что обладает законными основаниями для передачи сведений и материалов о третьих лицах, если они содержатся в обращении.")
    add_body(doc, "4.4. Обращение может быть направлено уполномоченным сотрудникам и компетентным государственным органам или органам местного самоуправления.")

    add_heading(doc, "5. Правила использования")
    add_body(doc, "5.1. Пользователь вправе использовать Чат-бот по назначению, получать доступную информацию о своём обращении и направлять вопросы по контактам Оператора.")
    add_body(doc, "5.2. Запрещается направлять заведомо ложные или незаконные материалы, нарушать права других лиц, вмешиваться в работу Чат-бота либо использовать его для массовой автоматизированной отправки сообщений.")
    add_body(doc, "5.3. Оператор обеспечивает работу Чат-бота и вправе ограничить доступ при злоупотреблении, а также изменять или временно приостанавливать работу Чат-бота для обновления и устранения неисправностей.")

    add_heading(doc, "6. Персональные данные и материалы")
    add_body(doc, "6.1. Обработка персональных данных регулируется отдельной Политикой и отдельным согласием Пользователя.")
    add_body(doc, "6.2. Непредоставление обязательных данных может сделать создание обращения невозможным. Фотографии, геолокация и точный адрес предоставляются только при наличии и необходимости.")
    add_body(doc, "6.3. Пользователь сохраняет права на свои материалы и разрешает использовать их только для рассмотрения и направления обращения.")

    add_heading(doc, "7. Работа Чат-бота и ответственность")
    add_body(doc, "7.1. В работе Чат-бота и мессенджера MAX возможны технические перерывы. Оператор принимает разумные меры для восстановления работы, но не отвечает за внешние системы вне своего контроля в пределах, допускаемых законом.")

    add_heading(doc, "8. Заключительные положения")
    add_body(doc, "8.1. Если Пользователь не обладает необходимым объёмом дееспособности, использование Чат-бота и предоставление данных осуществляются с участием законного представителя в предусмотренных законом случаях.")
    add_body(doc, "8.2. Применяется законодательство Российской Федерации. Споры разрешаются путём переговоров, а при недостижении соглашения — в установленном законом порядке.")
    add_body(doc, "8.3. Оператор может изменять Соглашение. Новая редакция применяется с указанной в ней даты и размещается в Чат-боте.")
    add_signature_block(doc)

    path = OUTPUT_DIR / "Пользовательское_соглашение_Искра_черновик.docx"
    doc.save(path)
    return path


def build_consent() -> Path:
    doc = configure_document("Согласие на обработку персональных данных", "согласие на обработку персональных данных")
    add_body(
        doc,
        "Пользователь чат-бота «Искра» в мессенджере MAX, действуя свободно, своей волей и в своём "
        "интересе, даёт конкретное, предметное, информированное, сознательное и однозначное согласие "
        "на обработку своих персональных данных на следующих условиях.",
    )

    add_heading(doc, "1. Оператор персональных данных")
    add_placeholder(doc, "Полное наименование или ФИО Оператора", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "ИНН", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "ОГРН/ОГРНИП", "[ЗАПОЛНИТЬ, ЕСЛИ ПРИМЕНИМО]")
    add_placeholder(doc, "Адрес Оператора", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "Электронная почта для обращений и отзыва согласия", "[ЗАПОЛНИТЬ ПОСЛЕ СОЗДАНИЯ]")

    add_heading(doc, "2. Цели обработки")
    for item in (
        "приём, регистрация, рассмотрение и направление обращения компетентным лицам и органам;",
        "связь с Пользователем, уточнение сведений, направление уведомлений и ответа;",
        "обеспечение работы и безопасности Чат-бота;",
        "формирование внутренней статистики и отчётности без распространения персональных данных.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "3. Перечень персональных данных")
    add_body(doc, "3.1. К персональным данным Пользователя относятся:")
    for item in (
        "фамилия, имя, отчество (при наличии);",
        "контактный номер телефона;",
        "сведения, автоматически передаваемые мессенджером MAX и необходимые для работы Чат-бота, включая идентификаторы Пользователя, чата и сообщений, дату и время взаимодействия;",
        "сведения о ходе и результате рассмотрения обращения.",
    ):
        add_bullet(doc, item)
    add_body(doc, "3.2. Материалами обращения являются муниципальный округ возникновения проблемы, текст обращения и предоставленные при необходимости фотографии, геолокация или точный адрес.")
    add_body(doc, "Эти материалы сами по себе не относятся к персональным данным Пользователя. Если они содержат сведения о конкретном или определяемом человеке, такие сведения обрабатываются как персональные данные.")
    add_body(doc, "Фотографии не используются Оператором для биометрической идентификации Пользователя.")

    add_heading(doc, "4. Действия и способы обработки")
    add_body(doc, "Оператор вправе осуществлять с указанными данными сбор, запись, систематизацию, накопление, хранение, уточнение, извлечение, использование, передачу (предоставление, доступ), обезличивание, блокирование, удаление и уничтожение.")
    add_body(doc, "Обработка осуществляется автоматизированным способом и без использования средств автоматизации. Распространение персональных данных неограниченному кругу лиц настоящим согласием не разрешается.")

    add_heading(doc, "5. Получатели данных")
    add_body(doc, "Доступ к данным в пределах указанных целей может предоставляться:")
    for item in (
        "уполномоченным сотрудникам Оператора;",
        "государственным органам и органам местного самоуправления, компетентным рассматривать обращение;",
        "оператору мессенджера MAX и иным лицам, обеспечивающим передачу сообщений и техническую работу Чат-бота, — только в необходимом объёме и на законном основании.",
    ):
        add_bullet(doc, item)
    add_body(doc, "Персональные данные хранятся и обрабатываются с использованием серверной инфраструктуры, расположенной на территории Российской Федерации.")

    add_heading(doc, "6. Срок обработки и хранения")
    add_body(doc, "Персональные данные и материалы обращения обрабатываются в течение срока рассмотрения обращения и 90 (девяноста) календарных дней после направления окончательного ответа, отклонения или иного завершения обращения, если более длительное хранение не требуется законодательством Российской Федерации.")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    r = set_run_font(p.add_run("[ДО ПУБЛИКАЦИИ: проверить на рабочем сервере автоматическое удаление данных и материалов обращения через 90 дней, а также регламент удаления их резервных копий.]"), bold=True)
    shade_run(r)

    add_heading(doc, "7. Отзыв согласия")
    add_body(doc, "Согласие может быть отозвано путём направления заявления Оператору:")
    add_placeholder(doc, "по электронной почте", "[АДРЕС ЭЛЕКТРОННОЙ ПОЧТЫ]")
    add_placeholder(doc, "или по почтовому адресу", "[АДРЕС ОПЕРАТОРА]")
    add_body(doc, "Заявление должно позволять идентифицировать Пользователя и содержать требование об отзыве согласия. После получения отзыва Оператор прекращает обработку и уничтожает данные в сроки, установленные законом, кроме случаев, когда обработка может быть продолжена на ином законном основании.")

    add_heading(doc, "8. Способ предоставления согласия")
    add_body(doc, "Согласие предоставляется отдельно от Пользовательского соглашения посредством нажатия Пользователем кнопки «Даю согласие на обработку персональных данных» до передачи обязательных данных.")
    add_body(doc, "Оператор фиксирует идентификатор Пользователя в MAX, дату и время предоставления согласия и редакцию документа для подтверждения факта его получения.")
    add_placeholder(doc, "Дата начала действия редакции", "[ДД.ММ.ГГГГ]")

    path = OUTPUT_DIR / "Согласие_на_обработку_ПДн_Искра_черновик.docx"
    doc.save(path)
    return path


def build_policy() -> Path:
    doc = configure_document("Политика обработки персональных данных", "политика обработки персональных данных")
    add_body(doc, "Настоящая Политика определяет порядок и условия обработки персональных данных физических лиц при использовании чат-бота «Искра» в мессенджере MAX.")

    add_heading(doc, "1. Общие положения")
    add_body(doc, "1.1. Политика разработана в соответствии с законодательством Российской Федерации о персональных данных и применяется ко всей информации, которую Оператор получает при работе Чат-бота.")
    add_body(doc, "1.2. До использования Чат-бота Пользователь знакомится с Политикой и Пользовательским соглашением и отдельно соглашается на обработку данных. Обработка данных самим мессенджером MAX регулируется его документами.")

    add_heading(doc, "2. Сведения об Операторе")
    add_placeholder(doc, "Полное наименование или ФИО", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "ИНН", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "ОГРН/ОГРНИП", "[ЗАПОЛНИТЬ, ЕСЛИ ПРИМЕНИМО]")
    add_placeholder(doc, "Адрес", "[ЗАПОЛНИТЬ]")
    add_placeholder(doc, "Электронная почта по вопросам персональных данных", "[ЗАПОЛНИТЬ ПОСЛЕ СОЗДАНИЯ]")

    add_heading(doc, "3. Категории субъектов и цели обработки")
    add_body(doc, "3.1. Обрабатываются данные Пользователей и сведения об иных лицах, указанные ими в обращениях.")
    add_body(doc, "3.2. Целями обработки являются:")
    for item in (
        "приём, рассмотрение и направление обращений компетентным лицам и органам;",
        "связь с Пользователем и направление ответа;",
        "обеспечение работы и безопасности Чат-бота;",
        "подготовка внутренней статистики и отчётности.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "4. Состав обрабатываемых данных и материалов")
    add_body(doc, "4.1. К персональным данным Пользователя относятся:")
    for item in (
        "ФИО Пользователя и контактный номер телефона;",
        "технические сведения, автоматически передаваемые мессенджером MAX и необходимые для работы Чат-бота;",
        "сведения о ходе и результате рассмотрения обращения.",
    ):
        add_bullet(doc, item)
    add_body(doc, "4.2. Материалами обращения являются округ возникновения проблемы, текст и предоставленные при необходимости фотографии, геолокация или точный адрес.")
    add_body(doc, "Эти материалы сами по себе не относятся к персональным данным Пользователя. Если они содержат сведения о конкретном или определяемом человеке, такие сведения обрабатываются как персональные данные.")
    add_body(doc, "Чат-бот не запрашивает округ проживания. Пользователь не должен сообщать избыточные сведения, не относящиеся к обращению.")

    add_heading(doc, "5. Правовые основания")
    add_body(doc, "Основаниями обработки являются согласие Пользователя, Пользовательское соглашение и применимые требования законодательства Российской Федерации.")

    add_heading(doc, "6. Порядок обработки и передачи")
    add_body(doc, "6.1. Обработка включает сбор, запись, систематизацию, накопление, хранение, уточнение, извлечение, использование, передачу, обезличивание, блокирование, удаление и уничтожение данных автоматизированным и неавтоматизированным способом.")
    add_body(doc, "6.2. Доступ предоставляется только в необходимом объёме уполномоченным сотрудникам, компетентным государственным органам и органам местного самоуправления.")
    add_body(doc, "6.3. Для передачи сообщений и размещения системы могут привлекаться технические подрядчики на законном основании. Данные не распространяются неограниченному кругу лиц, а фотографии не используются для биометрической идентификации.")

    add_heading(doc, "7. Локализация и хранение")
    add_body(doc, "7.1. Первичная запись и хранение персональных данных осуществляются с использованием серверной инфраструктуры, расположенной на территории Российской Федерации.")
    add_body(doc, "7.2. Персональные данные и материалы обращения хранятся в течение рассмотрения обращения и 90 (девяноста) календарных дней после его завершения, если законодательством не предусмотрен иной срок.")
    add_body(doc, "7.3. По окончании срока данные и подконтрольные Оператору копии удаляются или уничтожаются, если для хранения нет иного законного основания.")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    r = set_run_font(p.add_run("[ДО ПУБЛИКАЦИИ: проверить на рабочем сервере автоматическое удаление данных и материалов обращения через 90 дней, а также регламент удаления их резервных копий.]"), bold=True)
    shade_run(r)

    add_heading(doc, "8. Защита персональных данных")
    add_body(doc, "Оператор применяет необходимые правовые, организационные и технические меры защиты персональных данных в соответствии с законодательством Российской Федерации и с учётом актуальных угроз безопасности.")

    add_heading(doc, "9. Права Пользователя")
    add_body(doc, "Пользователь вправе получать сведения об обработке своих данных, требовать их уточнения, блокирования или уничтожения, отозвать согласие и обжаловать действия Оператора в порядке, установленном законом.")
    add_placeholder(doc, "Электронная почта для запросов", "[ЗАПОЛНИТЬ ПОСЛЕ СОЗДАНИЯ]")
    add_placeholder(doc, "Почтовый адрес", "[ЗАПОЛНИТЬ]")

    add_heading(doc, "10. Несовершеннолетние")
    add_body(doc, "Чат-бот не запрашивает возраст. Когда по закону требуется участие законного представителя, данные предоставляются с его участием.")

    add_heading(doc, "11. Изменение Политики")
    add_body(doc, "Оператор может обновлять Политику. Актуальная редакция и дата её вступления в силу размещаются в Чат-боте.")
    add_placeholder(doc, "Дата вступления в силу", "[ДД.ММ.ГГГГ]")

    path = OUTPUT_DIR / "Политика_обработки_ПДн_Искра_черновик.docx"
    doc.save(path)
    return path


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = [build_agreement(), build_consent(), build_policy()]
    for path in paths:
        print(path.resolve())


if __name__ == "__main__":
    main()

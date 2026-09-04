from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document


EXPECTED = {
    "Пользовательское_соглашение_Искра.docx": (
        "Принимаю пользовательское соглашение",
        "Согласие на обработку персональных данных запрашивается и фиксируется отдельно",
        "муниципальный округ, на территории которого находится указанная в обращении проблема",
    ),
    "Согласие_на_обработку_ПДн_Искра.docx": (
        "Даю согласие на обработку персональных данных",
        "90 (девяноста) календарных дней после направления окончательного ответа",
        "Распространение персональных данных неограниченному кругу лиц настоящим согласием не разрешается",
        "серверной инфраструктуры, расположенной на территории Российской Федерации",
        "Материалами обращения являются муниципальный округ возникновения проблемы, текст обращения",
        "Эти материалы сами по себе не относятся к персональным данным Пользователя",
    ),
    "Политика_обработки_ПДн_Искра.docx": (
        "Чат-бот не запрашивает округ проживания",
        "90 (девяноста) календарных дней после его завершения",
        "серверной инфраструктуры, расположенной на территории Российской Федерации",
        "Состав обрабатываемых данных и материалов",
        "Эти материалы сами по себе не относятся к персональным данным Пользователя",
    ),
}


def all_text(doc: Document) -> str:
    blocks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                blocks.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        blocks.extend(p.text for p in section.header.paragraphs)
        blocks.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(blocks)


def main() -> None:
    root = Path("legal")
    actual = {p.name for p in root.glob("*.docx")}
    assert actual == set(EXPECTED), f"Unexpected DOCX set: {sorted(actual)}"

    for name, required in EXPECTED.items():
        path = root / name
        doc = Document(path)
        text = all_text(doc)
        assert "Искра" in text and "MAX" in text
        assert "Сайта" not in text
        assert "Timeweb" not in text and "TimeWeb" not in text and "ТаймВэб" not in text
        assert "8.ользователь" not in text
        assert "ПРОЕКТ" not in text, f"Draft marker remains in {name}"
        assert "[" not in text and "]" not in text, f"Visible placeholders remain in {name}"
        assert "Министерство цифрового развития Калужской области" in text
        assert "ИНН: 4027138814" in text
        assert "1194027000221" in text and "11944027000221" not in text
        assert "min_digital@adm.kaluga.ru" in text
        assert "04.09.2026" in text
        for phrase in required:
            assert phrase in text, f"Missing phrase in {name}: {phrase}"

        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            assert "word/comments.xml" not in names, f"Unexpected comments in {name}"
            document_xml = archive.read("word/document.xml")
            assert b"<w:ins" not in document_xml and b"<w:del" not in document_xml

        print(f"OK {name}: paragraphs={len(doc.paragraphs)}, tables={len(doc.tables)}")


if __name__ == "__main__":
    main()

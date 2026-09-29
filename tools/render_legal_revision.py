"""Build current legal DOCX/PDF from legal/source without changing runtime settings."""
from __future__ import annotations
import hashlib
import json
import os
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak
from reportlab.lib.pagesizes import A4

ROOT = Path(__file__).resolve().parents[1]
LEGAL = ROOT / 'legal'
VERSION = '2.0'
DOCUMENTS = ('user-agreement', 'privacy-policy', 'personal-data-consent')
BRAND = 'На связи_регион40'

def blocks(name):
    return (LEGAL / 'source' / f'{name}.md').read_text(encoding='utf-8').strip().split('\n\n')

def build_docx(name):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(18)
    sec.left_margin = sec.right_margin = Mm(19)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_together = True
    for key, size in [('Title', 17), ('Heading 1', 12)]:
        style = doc.styles[key]
        style.font.name, style.font.size = 'Times New Roman', Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        for fonts in style.element.iter(qn('w:rFonts')):
            for attr in list(fonts.attrib):
                if 'theme' in attr.lower(): del fonts.attrib[attr]
    header = sec.header.paragraphs[0]
    header.text = f'{BRAND} · Документы сервиса'
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Страница ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    next_page = False
    for block in blocks(name):
        if block == '<!-- pagebreak -->':
            next_page = True
            continue
        elif block.startswith('# '): p = doc.add_paragraph(block[2:], 'Title')
        elif block.startswith('## '): p = doc.add_paragraph(block[3:], 'Heading 1')
        else:
            p = doc.add_paragraph(block)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if next_page:
            p.paragraph_format.page_break_before = True
            next_page = False
    doc.core_properties.title = blocks(name)[0][2:]
    doc.core_properties.author = 'Министерство цифрового развития Калужской области'
    doc.core_properties.subject = f'Чат-бот «{BRAND}», редакция {VERSION}'
    doc.save(LEGAL / f'{name}.docx')

def build_pdf(name):
    fonts = Path(os.environ.get('LEGAL_FONT_DIR', 'C:/Windows/Fonts'))
    for key, file in [('Legal', 'times.ttf'), ('LegalBold', 'timesbd.ttf')]:
        pdfmetrics.registerFont(TTFont(key, str(fonts / file)))
    styles = {
        'body': ParagraphStyle('body', fontName='Legal', fontSize=11, leading=13.2, spaceAfter=7, alignment=TA_JUSTIFY),
        'title': ParagraphStyle('title', fontName='LegalBold', fontSize=17, leading=20, spaceAfter=12, keepWithNext=True),
        'heading': ParagraphStyle('heading', fontName='LegalBold', fontSize=12, leading=14, spaceBefore=8, spaceAfter=6, keepWithNext=True),
    }
    story = []
    for block in blocks(name):
        if block == '<!-- pagebreak -->': story.append(PageBreak())
        elif block.startswith('# '): story.append(Paragraph(escape(block[2:]), styles['title']))
        elif block.startswith('## '): story.append(Paragraph(escape(block[3:]), styles['heading']))
        else: story.append(Paragraph(escape(block), styles['body']))
    def page(canvas, _doc):
        canvas.saveState()
        canvas.setFont('Legal', 8)
        canvas.drawString(54, A4[1]-32, f'{BRAND} · Документы сервиса')
        canvas.drawCentredString(A4[0]/2, 28, f'Страница {canvas.getPageNumber()}')
        canvas.restoreState()
    target = LEGAL / 'public' / f'{name}.pdf'
    SimpleDocTemplate(str(target), pagesize=A4, rightMargin=54, leftMargin=54,
        topMargin=51, bottomMargin=51, invariant=1, title=blocks(name)[0][2:], author='Министерство цифрового развития Калужской области').build(story, onFirstPage=page, onLaterPages=page)
    (LEGAL / 'pdf' / target.name).write_bytes(target.read_bytes())
    return hashlib.sha256(target.read_bytes()).hexdigest()

def main():
    for directory in ('public', 'pdf'): (LEGAL / directory).mkdir(exist_ok=True)
    result = {}
    for name in DOCUMENTS:
        build_docx(name)
        result[name] = {'version': VERSION, 'sha256': build_pdf(name)}
    (LEGAL / 'manifest.json').write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
    print('Built 3 DOCX and 3 PDF files; legal/manifest.json contains publication hashes.')

if __name__ == '__main__': main()

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
LEGAL_DIR = ROOT / "legal"


def replace_exact(document: Document, old: str, new: str, *previous_versions: str) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text == old or paragraph.text in previous_versions:
            paragraph.text = new
            return
        if paragraph.text == new:
            return
    raise RuntimeError(f"Paragraph not found: {old}")


def insert_before(paragraph: Paragraph, text: str) -> None:
    if paragraph._p.getprevious() is not None:
        previous = Paragraph(paragraph._p.getprevious(), paragraph._parent)
        if previous.text == text:
            return
    inserted = paragraph.insert_paragraph_before(text)
    inserted.style = paragraph.style


def compact_version_line(document: Document, date_text: str) -> None:
    version_text = "Редакция документа: 1.0"
    for paragraph in document.paragraphs:
        if paragraph.text != date_text:
            continue
        previous_element = paragraph._p.getprevious()
        if previous_element is not None:
            previous = Paragraph(previous_element, paragraph._parent)
            if previous.text == version_text:
                previous._element.getparent().remove(previous._element)
        if not paragraph.text.startswith(version_text):
            if paragraph.runs:
                paragraph.runs[0].text = f"{version_text}. {paragraph.runs[0].text}"
            else:
                paragraph.add_run(f"{version_text}. {date_text}")
        return
    combined = f"{version_text}. {date_text}"
    if any(paragraph.text == combined for paragraph in document.paragraphs):
        return
    raise RuntimeError(f"Effective date paragraph not found: {date_text}")


def update_policy() -> None:
    path = LEGAL_DIR / "Политика_обработки_ПДн_Искра_черновик.docx"
    document = Document(path)
    replace_exact(
        document,
        "1.2. До использования Чат-бота Пользователь знакомится с Политикой и Пользовательским соглашением и отдельно соглашается на обработку данных. Обработка данных самим мессенджером MAX регулируется его документами.",
        "1.2. Перед первым созданием обращения Пользователь знакомится с Политикой, Пользовательским соглашением и Согласием на обработку персональных данных. Соглашение и Согласие подтверждаются разными кнопками. Обработка данных самим мессенджером MAX регулируется его документами.",
        "1.2. Перед первым созданием обращения Пользователю предоставляется доступ к настоящей Политике, Пользовательскому соглашению и Согласию на обработку персональных данных. Принятие Пользовательского соглашения и предоставление согласия на обработку персональных данных подтверждаются отдельными действиями в Чат-боте. Обработка данных самим мессенджером MAX регулируется его документами.",
    )
    replace_exact(
        document,
        "[ДО ПУБЛИКАЦИИ: проверить на рабочем сервере автоматическое удаление данных и материалов обращения через 90 дней, а также регламент удаления их резервных копий.]",
        "Резервные копии с данными и материалами обращений хранятся не более 90 календарных дней, если иной срок не требуется законодательством Российской Федерации.",
        "Резервные копии, содержащие данные и материалы обращений, удаляются по установленной Оператором политике хранения в срок не более 90 календарных дней, если более длительное хранение не требуется законодательством Российской Федерации.",
    )
    compact_version_line(document, "Дата вступления в силу: [ДД.ММ.ГГГГ]")
    for paragraph in document.paragraphs:
        if paragraph.text == "11. Изменение Политики":
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(2)
        elif paragraph.text.startswith("Оператор может обновлять Политику."):
            paragraph.paragraph_format.space_after = Pt(0)
    document.save(path)


def update_agreement() -> None:
    path = LEGAL_DIR / "Пользовательское_соглашение_Искра_черновик.docx"
    document = Document(path)
    replace_exact(
        document,
        "3.1. До начала использования Чат-бота Пользователь знакомится с настоящим Соглашением и Политикой обработки персональных данных.",
        "3.1. Перед первым созданием обращения Пользователю предоставляется доступ к настоящему Соглашению, Политике обработки персональных данных и отдельному Согласию на обработку персональных данных.",
    )
    replace_exact(
        document,
        "8.3. Оператор может изменять Соглашение. Новая редакция применяется с указанной в ней даты и размещается в Чат-боте.",
        "8.3. Оператор может изменять Соглашение. Новая редакция применяется с указанной в ней даты и размещается в Чат-боте. Если изменения требуют нового подтверждения, до создания следующего обращения Пользователю предлагается принять новую редакцию отдельным действием.",
    )
    compact_version_line(document, "Дата вступления в силу: [ДД.ММ.ГГГГ]")
    document.save(path)


def update_consent() -> None:
    path = LEGAL_DIR / "Согласие_на_обработку_ПДн_Искра_черновик.docx"
    document = Document(path)
    replace_exact(
        document,
        "[ДО ПУБЛИКАЦИИ: проверить на рабочем сервере автоматическое удаление данных и материалов обращения через 90 дней, а также регламент удаления их резервных копий.]",
        "Резервные копии, содержащие персональные данные и материалы обращений, удаляются по установленной Оператором политике хранения в срок не более 90 календарных дней, если более длительное хранение не требуется законодательством Российской Федерации.",
    )
    replace_exact(
        document,
        "Согласие предоставляется отдельно от Пользовательского соглашения посредством нажатия Пользователем кнопки «Даю согласие на обработку персональных данных» до передачи обязательных данных.",
        "Согласие предоставляется отдельно от Пользовательского соглашения посредством нажатия Пользователем кнопки «Даю согласие на обработку персональных данных» до передачи обязательных данных. Нажатие кнопки «Принимаю пользовательское соглашение» не считается предоставлением настоящего согласия.",
    )
    replace_exact(
        document,
        "Оператор фиксирует идентификатор Пользователя в MAX, дату и время предоставления согласия и редакцию документа для подтверждения факта его получения.",
        "Оператор фиксирует идентификатор Пользователя в MAX, дату и время предоставления согласия, редакцию и контрольную сумму документа, а также идентификатор подтверждающего действия в MAX для подтверждения факта получения согласия.",
    )
    compact_version_line(document, "Дата начала действия редакции: [ДД.ММ.ГГГГ]")
    document.save(path)


def main() -> None:
    update_policy()
    update_agreement()
    update_consent()
    print("Updated 3 legal DOCX files")


if __name__ == "__main__":
    main()
